import os
import re

import streamlit as st
from dotenv import load_dotenv
from google import genai
from supabase import create_client, Client

from io import BytesIO
from datetime import datetime

from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph
from docx import Document
from supabase import create_client

# -----------------------------
# Load Environment
# -----------------------------
load_dotenv()

API_KEY = os.getenv("GOOGLE_API_KEY")
if not API_KEY:
    st.error("❌ GOOGLE_API_KEY not found.")
    st.stop()

client = genai.Client(api_key=API_KEY)
supabase_url = os.getenv("SUPABASE_URL")
supabase_key = os.getenv("SUPABASE_PUBLISHABLE_KEY")

if not supabase_url or not supabase_key:
    st.error("❌ Supabase credentials not found.")
    st.stop()

supabase = create_client(
    supabase_url,
    supabase_key,
)

# -----------------------------
# Page Config
# -----------------------------

st.set_page_config(
    page_title="StudyGuide AI",
    page_icon="📚",
    layout="wide"
)

# -----------------------------
# Custom UI Theme
# -----------------------------

st.markdown(
    """
    <style>

    .stApp {
        background-color: #0E1117;
    }

    .main .block-container {
        padding-top: 2rem;
        padding-bottom: 2rem;
        max-width: 1100px;
    }

    h1, h2, h3 {
        color: #FAFAFA;
    }

    .stButton > button {
        width: 100%;
        border-radius: 12px;
        height: 3em;
        font-weight: bold;
    }

    .stDownloadButton > button {
        width: 100%;
        border-radius: 12px;
        height: 3em;
        font-weight: bold;
    }

    textarea {
        border-radius: 10px !important;
    }

    </style>
    """,
    unsafe_allow_html=True,
)



# -----------------------------
# Session
# -----------------------------

if "history" not in st.session_state:
    st.session_state.history = []
if "user" not in st.session_state:
    st.session_state.user = None

if "chat_count" not in st.session_state:
    st.session_state.chat_count = 0

# -----------------------------
# Login Gate
# -----------------------------

if st.session_state.user is None:

    st.title("🔐 StudyGuide AI Login")
    st.caption("Sign in to access your AI study assistant.")

    login_email = st.text_input(
        "Email",
        placeholder="student@example.com",
    )

    login_password = st.text_input(
        "Password",
        type="password",
    )

    if st.button("Login", type="primary"):

        if not login_email or not login_password:
            st.warning("Please enter both email and password.")

        else:
            try:
                login_response = supabase.auth.sign_in_with_password(
                    {
                        "email": login_email,
                        "password": login_password,
                    }
                )

                st.session_state.user = login_response.user
                st.success("Login successful.")
                st.rerun()

            except Exception as e:
                st.error(f"Login failed: {e}")

    st.markdown("---")
    st.subheader("New here?")

    signup_email = st.text_input(
        "New Email",
        key="signup_email",
    )

    signup_password = st.text_input(
        "New Password",
        type="password",
        key="signup_password",
    )

    if st.button("Create Account"):

        if not signup_email or not signup_password:
            st.warning("Please enter email and password.")

        else:
            try:
                supabase.auth.sign_up(
                    {
                        "email": signup_email,
                        "password": signup_password,
                        "options": {
                            "email_redirect_to": "http://localhost:8501",
                        },
                    }
                )

                st.success(
                    "Account created. Check your email for confirmation, then log in."
                )

            except Exception as e:
                st.error(f"Sign up failed: {e}")

    st.stop()
# -----------------------------
# Sidebar
# -----------------------------

with st.sidebar:

    st.title("📚 StudyGuide AI")

    st.markdown("---")

    department = st.selectbox(
        "Department",
        [
            "EEE",
            "CSE",
            "Physics",
            "Chemistry",
            "Mathematics",
            "Biology",
            "Others"
        ]
    )

    subject = st.text_input(
        "Course / Subject",
        placeholder="Control System"
    )

    topic = st.text_input(
        "Topic",
        placeholder="Routh Hurwitz Stability"
    )

    mode = st.selectbox(
        "Study Mode",
        [
            "Study Plan",
            "Short Notes",
            "Quiz",
            "Viva Questions",
            "Concept Explanation"
        ]
    )

    language = st.selectbox(
        "Language",
        [
            "English",
            "Bangla",
            "Bangla + English"
        ]
    )

    difficulty = st.selectbox(
        "Difficulty",
        [
            "Easy",
            "Medium",
            "Hard"
        ]
    )

    days = st.slider(
        "Study Duration (Days)",
        1,
        30,
        7
    )

# -----------------------------
# Title
# -----------------------------

st.title("📚 StudyGuide AI")

st.caption(
    "AI Tutor for University Students | Study Plans | Viva | Quiz | Notes | Exam Preparation"
)

st.markdown("---")

# -----------------------------
# Security Guardrails
# -----------------------------

EMAIL_REGEX = re.compile(
    r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+"
)

PHONE_REGEX = re.compile(
    r"\b\d{3}[-.]?\d{3}[-.]?\d{4}\b"
)

INJECTION_KEYWORDS = [

    "ignore previous instructions",

    "system prompt",

    "override rules",

    "you are now a",

    "bypass restrictions",

]

CHEATING_KEYWORDS = [

    "do my homework",

    "write my assignment",

    "write my essay",

    "solve my exam",

    "give me exam answers",

    "cheat sheet",

]

def security_check(text):

    clean = EMAIL_REGEX.sub("[EMAIL]", text)

    clean = PHONE_REGEX.sub("[PHONE]", clean)

    lower = clean.lower()

    if any(x in lower for x in INJECTION_KEYWORDS):

        return False, "Prompt Injection detected."

    if any(x in lower for x in CHEATING_KEYWORDS):

        return False, "Academic Integrity violation detected."

    return True, clean

def create_pdf(question, answer):

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()

    story = []

    story.append(Paragraph("<b>StudyGuide AI Report</b>", styles["Title"]))

    story.append(
        Paragraph(
            f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}",
            styles["Normal"],
        )
    )

    story.append(Paragraph("<br/><b>Question</b>", styles["Heading2"]))
    story.append(Paragraph(question, styles["BodyText"]))

    story.append(Paragraph("<br/><b>Answer</b>", styles["Heading2"]))
    story.append(Paragraph(answer.replace("\n", "<br/>"), styles["BodyText"]))

    doc.build(story)

    buffer.seek(0)

    return buffer


# ==============================
# এখানে নতুন function paste হবে
# ==============================

def create_history_pdf(history):
    buffer = BytesIO()

    doc = SimpleDocTemplate(buffer)
    styles = getSampleStyleSheet()

    story = []

    story.append(
        Paragraph(
            "<b>StudyGuide AI - Conversation History</b>",
            styles["Title"],
        )
    )

    story.append(
        Paragraph(
            f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}",
            styles["Normal"],
        )
    )

    for item in history:
        story.append(
            Paragraph(
                f"<b>Conversation #{item['id']}</b>",
                styles["Heading2"],
            )
        )

        story.append(
            Paragraph(
                "<b>Question:</b>",
                styles["Heading3"],
            )
        )

        story.append(
            Paragraph(
                item["question"],
                styles["BodyText"],
            )
        )

        story.append(
            Paragraph(
                "<b>Answer:</b>",
                styles["Heading3"],
            )
        )

        story.append(
            Paragraph(
                item["answer"].replace("\n", "<br/>"),
                styles["BodyText"],
            )
        )

    doc.build(story)

    buffer.seek(0)

    return buffer
    ...

def create_docx(question, answer):
    document = Document()

    document.add_heading("StudyGuide AI Report", level=1)

    document.add_paragraph(
        f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}"
    )

    document.add_heading("Question", level=2)
    document.add_paragraph(question)

    document.add_heading("Answer", level=2)
    document.add_paragraph(answer)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer

def type_animation(text): 
    placeholder = st.empty() 
    displayed_text = "" 
    for char in text:
            displayed_text += char
            placeholder.markdown(displayed_text) 
    return displayed_text


st.markdown("## Ask StudyGuide AI")

user_prompt = st.text_area(

    "Enter your request",

    height=180,

    placeholder="""
Examples:

Create a 7-day study plan for Digital Electronics.

Generate viva questions for Control System.

Explain Kirchhoff's Voltage Law.

Prepare a quiz on Thermodynamics.

Create short notes on Organic Chemistry.
"""

)

generate = st.button("Generate")


# -----------------------------
# Generate Response
# -----------------------------

if generate:

    valid, result = security_check(user_prompt)

    if not valid:
        st.error(result)
        st.stop()

    prompt = f"""
You are StudyGuide AI.

You are an academic tutor for university and school students.

Department:
{department}

Course:
{subject}

Topic:
{topic}

Study Mode:
{mode}

Difficulty:
{difficulty}

Study Duration:
{days} days

Language:
{language}

Student Request:
{result}

Instructions:

- Explain clearly.
- Be exam focused.
- Use bullet points.
- Give practical examples where possible.
- If Study Mode is Study Plan, create a day-by-day plan.
- If Study Mode is Quiz, generate 10 questions with answers.
- If Study Mode is Viva Questions, generate 15 viva questions with short answers.
- If Study Mode is Short Notes, create revision notes.
- If Study Mode is Concept Explanation, explain step by step.
- If language is Bangla + English, mix both naturally.
- End with a short summary.
- Keep the answer neat and easy to read.
- Never help with cheating or exam misconduct.
"""
    
    try:    
        with st.spinner("🤖 StudyGuide AI is preparing your answer..."):    

            response = client.models.generate_content(    
            model="gemini-2.5-flash",    
            contents=prompt,    
        )    

        answer = response.text    

    except Exception as e:    
        st.error(f"⚠️ Error: {e}")    
        st.stop()


    st.session_state.chat_count += 1

    st.session_state.history.append(
    {
        "id": st.session_state.chat_count,
        "department": department,
        "subject": subject,
        "topic": topic,
        "mode": mode,
        "language": language,
        "difficulty": difficulty,
        "question": user_prompt,
        "answer": answer,
        "timestamp": datetime.now().strftime("%d-%m-%Y %H:%M:%S"),
    }
)
    try:
        supabase.table("study_history").insert(
        {
            "department": department,
            "subject": subject,
            "topic": topic,
            "mode": mode,
            "language": language,
            "difficulty": difficulty,
            "question": user_prompt,
            "answer": answer,
        }
    ).execute()

    except Exception as e:
        st.warning(f"Supabase save failed: {e}")
# -----------------------------
# History Controls
# -----------------------------

if st.session_state.history:

    col1, col2 = st.columns([5, 1])

    with col2:
        if st.button("🗑️ Clear History"):
            st.session_state.history = []
            st.session_state.chat_count = 0
            st.rerun()
# -----------------------------
# Chat History
# -----------------------------

if st.session_state.history:

    st.markdown("---")
    st.header(f"💬 Conversation ({len(st.session_state.history)})")
    
    total_conversations = len(st.session_state.history)

    average_answer_length = (
    sum(len(item["answer"]) for item in st.session_state.history)
    // total_conversations
    if total_conversations > 0
    else 0
)

    st.info(
    f"""
    📊 Statistics

    • Total Conversations: {total_conversations}

    • Average Answer Length: {average_answer_length} characters
    """
    )

    search_query = st.text_input(
        "🔍 Search Conversations",
        placeholder="Search by question, answer, topic..."
    )

for item in reversed(st.session_state.history):

    with st.expander(
        f"💬 Conversation #{item['id']} - {item['topic']}"
    ):
        st.caption(
                f"🕒 {item.get('timestamp', 'No timestamp')}"
            )


        st.markdown("### 👤 Question")
        st.write(item["question"])

        st.markdown("### 🤖 StudyGuide AI")
        type_animation(item["answer"])


        if st.button(
            f"🗑️ Delete Conversation #{item['id']}",
            key=f"delete_{item['id']}"
        ):
            st.session_state.history = [
                chat for chat in st.session_state.history
                if chat["id"] != item["id"]
            ]
            st.rerun()



# -----------------------------
# Download
# -----------------------------

    if st.session_state.history:
        
        
        latest = st.session_state.history[-1]["answer"]

        st.download_button(
            "Download Latest Answer",
            latest,
            file_name="studyguide_answer.txt",
            key="download_latest_txt",
        )

        pdf_buffer = create_pdf(
            st.session_state.history[-1]["question"],
            st.session_state.history[-1]["answer"],
        )

        st.download_button(
            "📄 Download as PDF",
            data=pdf_buffer,
            file_name="studyguide_answer.pdf",
            mime="application/pdf",
            key="download_latest_pdf",
        )

        docx_buffer = create_docx(
            st.session_state.history[-1]["question"],
            st.session_state.history[-1]["answer"],
        )

        st.download_button(
            "📝 Download as DOCX",
            data=docx_buffer,
            file_name="studyguide_answer.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key="download_latest_docx",
        )


        markdown_content = f"""# StudyGuide AI Report

        **Question:**

        {st.session_state.history[-1]["question"]}

        ---

        **Answer:**

        {st.session_state.history[-1]["answer"]}
        """

        st.download_button(
            "⬇️ Download as Markdown",
            data=markdown_content,
            file_name="studyguide_answer.md",
            mime="text/markdown",
        )
        
        if st.button("📋 Copy Latest Answer"):
            st.code(
                st.session_state.history[-1]["answer"],
                language=None,
            )
            st.success("Answer is ready. Select the text above and press Ctrl+C.")
        history_pdf = create_history_pdf(
            st.session_state.history
        )

        st.download_button(
            "📚 Export Full History as PDF",
            data=history_pdf,
            file_name="studyguide_history.pdf",
            mime="application/pdf",
        )





